import os
import json
import threading

os.environ.setdefault("OPENAI_API_KEY", "test-key")
os.environ.setdefault("DATABASE_URL", "sqlite:///:memory:")
os.environ.setdefault("JWT_SECRET", "test-secret-that-is-at-least-32-bytes-long")

from fastapi import WebSocketDisconnect
from fastapi.testclient import TestClient
from io import BytesIO
from types import SimpleNamespace

import pytest
from docx import Document

from app.main import app, get_db
from app.document_loader import extract_document_text
import app.main as main
import app.ai as ai
import app.agent as agent
import app.crud as crud
from app.security import decode_access_token, hash_password, verify_password

client = TestClient(app)


@pytest.fixture
def authenticated_user():
    app.dependency_overrides[main.get_current_user] = lambda: SimpleNamespace(
        id=1,
        name="Prabhati",
        email="prabhati@example.com",
    )
    yield
    app.dependency_overrides.pop(main.get_current_user, None)


def test_health_check():
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_chat_ui_and_assets_are_served():
    page_response = client.get("/")
    asset_response = client.get("/static/chat.js")

    assert page_response.status_code == 200
    assert "Make your documents" in page_response.text
    assert asset_response.status_code == 200
    assert "documents/upload" in asset_response.text
    assert page_response.headers["cache-control"] == "no-cache"
    assert asset_response.headers["cache-control"] == "no-cache"


def test_list_documents_requires_authentication():
    response = client.get("/documents")

    assert response.status_code == 401


def test_list_documents_returns_current_users_documents(monkeypatch, authenticated_user):
    queried_user_ids = []

    def fake_list_documents(db, user_id):
        queried_user_ids.append(user_id)
        return [SimpleNamespace(id=4, title="My CV"), SimpleNamespace(id=9, title="Notes")]

    monkeypatch.setattr(crud, "list_documents", fake_list_documents)

    response = client.get("/documents")

    assert response.status_code == 200
    assert response.json() == [{"id": 4, "title": "My CV"}, {"id": 9, "title": "Notes"}]
    assert queried_user_ids == [1]


def test_register_creates_user_and_returns_jwt(monkeypatch):
    monkeypatch.setattr(crud, "get_user_by_name", lambda db, name: None)
    monkeypatch.setattr(crud, "get_user_by_email", lambda db, email: None)
    monkeypatch.setattr(
        crud,
        "create_user",
        lambda db, user: SimpleNamespace(id=12, name=user.name, email=user.email),
    )
    response = client.post(
        "/auth/register",
        json={"name": "Prabhati", "email": "prabhati@example.com", "password": "secret123"},
    )

    assert response.status_code == 201
    assert response.json()["user"]["name"] == "Prabhati"
    assert decode_access_token(response.json()["access_token"]) == 12


def test_register_rejects_duplicate_email(monkeypatch):
    monkeypatch.setattr(crud, "get_user_by_name", lambda db, name: None)
    monkeypatch.setattr(crud, "get_user_by_email", lambda db, email: object())

    response = client.post(
        "/auth/register",
        json={"name": "Prabhati", "email": "prabhati@example.com", "password": "secret123"},
    )

    assert response.status_code == 409


def test_login_returns_jwt(monkeypatch):
    user = SimpleNamespace(
        id=12,
        name="Prabhati",
        email="prabhati@example.com",
        password_hash=hash_password("secret123"),
    )
    monkeypatch.setattr(crud, "get_user_by_email", lambda db, email: user)

    response = client.post(
        "/auth/login",
        json={"email": "prabhati@example.com", "password": "secret123"},
    )
    assert response.status_code == 200
    assert response.json()["user"]["name"] == "Prabhati"
    assert decode_access_token(response.json()["access_token"]) == 12


def test_current_user_endpoint_reads_user_from_valid_jwt(monkeypatch):
    user = SimpleNamespace(id=12, name="Prabhati", email="prabhati@example.com")
    monkeypatch.setattr(crud, "get_user", lambda db, user_id: user if user_id == 12 else None)
    token = main.create_access_token(12)

    response = client.get("/auth/me", headers={"Authorization": f"Bearer {token}"})

    assert response.status_code == 200
    assert response.json()["name"] == "Prabhati"


def test_current_user_endpoint_rejects_invalid_jwt():
    response = client.get("/auth/me", headers={"Authorization": "Bearer invalid-token"})

    assert response.status_code == 401


def test_ollama_embedding_uses_local_embedding_api(monkeypatch):
    import app.ai as ai

    captured = {}

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def read(self):
            return json.dumps({"embeddings": [[0.25] * 768]}).encode()

    def fake_urlopen(request, timeout):
        captured["url"] = request.full_url
        captured["payload"] = json.loads(request.data)
        captured["timeout"] = timeout
        return FakeResponse()

    monkeypatch.setattr(ai, "EMBEDDING_PROVIDER", "ollama")
    monkeypatch.setattr(ai, "EMBEDDING_DIMENSIONS", 768)
    monkeypatch.setattr(ai, "urlopen", fake_urlopen)

    embedding = ai.generate_embedding("CV skill")

    assert len(embedding) == 768
    assert captured["url"].endswith("/api/embed")
    assert captured["payload"]["model"] == ai.OLLAMA_EMBEDDING_MODEL
    assert captured["payload"]["input"] == ["CV skill"]


def test_generate_embeddings_sends_chunks_in_batches(monkeypatch):
    batches = []

    def fake_embed_batch(texts):
        batches.append(len(texts))
        return [[0.0] for _ in texts]

    monkeypatch.setattr(ai, "_embed_batch", fake_embed_batch)
    monkeypatch.setattr(ai, "EMBEDDING_BATCH_SIZE", 16)

    embeddings = ai.generate_embeddings(["chunk"] * 40)

    assert batches == [16, 16, 8]
    assert len(embeddings) == 40


def test_ollama_chat_uses_local_chat_api(monkeypatch):
    import app.ai as ai

    captured = {}

    class FakeResponse:
        def __enter__(self):
            return self

        def __exit__(self, *args):
            return None

        def read(self):
            return json.dumps({"message": {"content": "A local answer."}}).encode()

    def fake_urlopen(request, timeout):
        captured["url"] = request.full_url
        captured["payload"] = json.loads(request.data)
        return FakeResponse()

    monkeypatch.setattr(ai, "AI_PROVIDER", "ollama")
    monkeypatch.setattr(ai, "urlopen", fake_urlopen)

    history = [
        {"role": "user", "content": "Summarize my experience."},
        {"role": "assistant", "content": "Earlier answer."},
    ]
    answer = ai.generate_response("Explain the first project.", context="Python", history=history)

    assert answer == "A local answer."
    assert captured["url"].endswith("/api/chat")
    assert captured["payload"]["model"] == ai.OLLAMA_CHAT_MODEL
    assert captured["payload"]["stream"] is False
    assert "Python" not in captured["payload"]["messages"][0]["content"]
    assert captured["payload"]["messages"][1:-1] == history
    last_message = captured["payload"]["messages"][-1]
    assert last_message["role"] == "user"
    assert "Python" in last_message["content"]
    assert last_message["content"].endswith("Question: Explain the first project.")
    assert "must never override the document context" in captured["payload"]["messages"][0]["content"]


def test_ollama_request_reports_timeout(monkeypatch):
    def raise_timeout(request, timeout):
        raise TimeoutError

    monkeypatch.setattr(ai, "urlopen", raise_timeout)

    with pytest.raises(ai.OllamaTimeoutError, match="did not respond within 120 seconds"):
        ai._ollama_request("/api/chat", {})


def test_ollama_stream_yields_each_piece_of_the_answer(monkeypatch):
    lines = [
        json.dumps({"message": {"content": "Hello"}}).encode() + b"\n",
        json.dumps({"message": {"content": " world"}}).encode() + b"\n",
        json.dumps({"message": {"content": ""}, "done": True}).encode() + b"\n",
    ]
    captured = {}

    class FakeResponse:
        def __enter__(self):
            return iter(lines)

        def __exit__(self, *args):
            return None

    def fake_urlopen(request, timeout):
        captured["payload"] = json.loads(request.data)
        return FakeResponse()

    monkeypatch.setattr(ai, "AI_PROVIDER", "ollama")
    monkeypatch.setattr(ai, "urlopen", fake_urlopen)

    assert list(ai.stream_response("Hi", context="CV")) == ["Hello", " world"]
    assert captured["payload"]["stream"] is True


def test_ollama_stream_raises_when_ollama_reports_an_error(monkeypatch):
    class FakeResponse:
        def __enter__(self):
            return iter([json.dumps({"error": "out of memory"}).encode() + b"\n"])

        def __exit__(self, *args):
            return None

    monkeypatch.setattr(ai, "AI_PROVIDER", "ollama")
    monkeypatch.setattr(ai, "urlopen", lambda request, timeout: FakeResponse())

    with pytest.raises(RuntimeError, match="out of memory"):
        list(ai.stream_response("Hi", context="CV"))


def test_agent_uses_plain_rag_for_ollama(monkeypatch):
    monkeypatch.setattr(ai, "AI_PROVIDER", "ollama")
    monkeypatch.setattr(agent, "_search_knowledge_base", lambda db, user_id, query: "CV context")
    monkeypatch.setattr(ai, "generate_response", lambda prompt, context=None: f"{context}: {prompt}")

    response = agent.run_agent(None, "Summarize my CV", user_id=7)

    assert response == "CV context: Summarize my CV"


def test_reindex_endpoint_rebuilds_only_authenticated_user_chunks(monkeypatch, authenticated_user):
    called = []

    def fake_reindex(db, user_id):
        called.append(user_id)
        return 3

    monkeypatch.setattr(crud, "reembed_user_document_chunks", fake_reindex)

    response = client.post("/documents/reindex")

    assert response.status_code == 200
    assert response.json() == {"reindexed_chunks": 3}
    assert called == [1]


def test_reembed_user_document_chunks_generates_embeddings(monkeypatch):
    from app import models

    chunks = [(11, "First CV section"), (12, "Second CV section")]
    updated = []

    class Query:
        def join(self, model):
            return self

        def filter(self, condition):
            return self

        def all(self):
            return chunks

        def update(self, values, synchronize_session):
            updated.append((self.chunk_id, values, synchronize_session))
            return 1

        @property
        def chunk_id(self):
            return len(updated) + 11

    class FakeDB:
        committed = False

        def query(self, *query_models):
            if len(query_models) == 2:
                assert query_models[0] is models.DocumentChunk.id
                assert query_models[1] is models.DocumentChunk.content
            else:
                assert query_models == (models.DocumentChunk,)
            return Query()

        def commit(self):
            self.committed = True

        def rollback(self):
            pytest.fail("Rollback should not happen")

    monkeypatch.setattr(
        ai,
        "generate_embeddings",
        lambda contents: [[float(len(content))] * 768 for content in contents],
    )
    db = FakeDB()

    count = crud.reembed_user_document_chunks(db, user_id=1)

    assert count == 2
    assert len(updated) == 2
    assert updated[0][1][models.DocumentChunk.embedding] == [float(len("First CV section"))] * 768
    assert updated[1][1][models.DocumentChunk.embedding] == [float(len("Second CV section"))] * 768
    assert all(sync is False for _, _, sync in updated)
    assert db.committed


def test_password_hash_is_verified_without_storing_plaintext():
    password_hash = hash_password("secret123")

    assert password_hash != "secret123"
    assert verify_password("secret123", password_hash)
    assert not verify_password("wrong-password", password_hash)


@pytest.fixture
def signed_in_socket(monkeypatch):
    monkeypatch.setattr(main, "_socket_user_id", lambda token: 1)
    saved = {}
    monkeypatch.setattr(crud, "create_chat_session", lambda db, user_id, title: SimpleNamespace(id=5))

    def fake_create_chat_message(db, user_id, prompt, response, session_id=None):
        saved.update(user_id=user_id, reply=response, session_id=session_id)

    monkeypatch.setattr(crud, "create_chat_message", fake_create_chat_message)
    with client.websocket_connect("/ws/chat") as websocket:
        websocket.send_json({"type": "auth", "token": "valid"})
        assert websocket.receive_json() == {"type": "ready"}
        yield websocket, saved


def test_chat_socket_rejects_invalid_token():
    with client.websocket_connect("/ws/chat") as websocket:
        websocket.send_json({"type": "auth", "token": "invalid-token"})
        with pytest.raises(WebSocketDisconnect) as closed:
            websocket.receive_json()

    assert closed.value.code == main.SOCKET_UNAUTHORIZED


def test_chat_socket_streams_tokens_and_saves_the_answer(monkeypatch, signed_in_socket):
    websocket, saved = signed_in_socket
    monkeypatch.setattr(main, "_prepare_socket_answer", lambda user_id, message: ([], "CV", ["My CV"]))
    monkeypatch.setattr(
        ai,
        "stream_response",
        lambda prompt, context=None, history=None: iter(["Hello ", "world"]),
    )

    websocket.send_json({"type": "question", "prompt": "Hi"})

    assert websocket.receive_json() == {"type": "sources", "titles": ["My CV"]}
    assert websocket.receive_json() == {"type": "token", "text": "Hello "}
    assert websocket.receive_json() == {"type": "token", "text": "world"}
    assert websocket.receive_json() == {"type": "done", "session_id": 5}
    assert saved == {"user_id": 1, "reply": "Hello world", "session_id": 5}


def test_chat_socket_stop_keeps_the_partial_answer(monkeypatch, signed_in_socket):
    websocket, saved = signed_in_socket
    release = threading.Event()

    def slow_stream(prompt, context=None, history=None):
        yield "Partial "
        release.wait(5)
        yield "never sent"

    monkeypatch.setattr(main, "_prepare_socket_answer", lambda user_id, message: ([], "CV", []))
    monkeypatch.setattr(ai, "stream_response", slow_stream)

    websocket.send_json({"type": "question", "prompt": "Hi"})
    assert websocket.receive_json()["type"] == "sources"
    assert websocket.receive_json() == {"type": "token", "text": "Partial "}

    websocket.send_json({"type": "stop"})

    assert websocket.receive_json() == {"type": "stopped", "session_id": 5}
    assert saved["reply"] == "Partial "
    release.set()


def test_chat_socket_refuses_another_users_conversation(monkeypatch, signed_in_socket):
    websocket, saved = signed_in_socket
    monkeypatch.setattr(crud, "get_chat_session", lambda db, session_id, user_id: None)
    monkeypatch.setattr(
        ai,
        "generate_embedding",
        lambda prompt: pytest.fail("Search should not run for someone else's conversation"),
    )

    websocket.send_json({"type": "question", "prompt": "Hi", "session_id": 99})

    assert websocket.receive_json() == {
        "type": "error",
        "message": "That conversation no longer exists.",
    }
    assert saved == {}


def test_chat_socket_reports_model_failure_without_saving(monkeypatch, signed_in_socket):
    websocket, saved = signed_in_socket

    def failed_prepare(user_id, message):
        raise RuntimeError("Could not reach Ollama")

    monkeypatch.setattr(main, "_prepare_socket_answer", failed_prepare)

    websocket.send_json({"type": "question", "prompt": "Hi"})

    assert websocket.receive_json() == {"type": "error", "message": "Could not reach Ollama"}
    assert saved == {}


def test_session_messages_hidden_from_other_users(monkeypatch, authenticated_user):
    monkeypatch.setattr(crud, "get_chat_session", lambda db, session_id, user_id: None)

    assert client.get("/sessions/99/messages").status_code == 404
    assert client.delete("/sessions/99").status_code == 404


def test_list_sessions_returns_current_users_conversations(monkeypatch, authenticated_user):
    queried = []

    def fake_list_chat_sessions(db, user_id):
        queried.append(user_id)
        return [SimpleNamespace(id=3, title="About my CV")]

    monkeypatch.setattr(crud, "list_chat_sessions", fake_list_chat_sessions)

    response = client.get("/sessions")

    assert response.json() == [{"id": 3, "title": "About my CV"}]
    assert queried == [1]


def test_chat_requires_authentication():
    response = client.post("/chat", json={"prompt": "Hi"})

    assert response.status_code == 401


def test_chat_uses_authenticated_user(monkeypatch, authenticated_user):
    history = [
        {"role": "user", "content": "Summarize my experience."},
        {"role": "assistant", "content": "Earlier answer."},
    ]
    monkeypatch.setattr(crud, "get_recent_chat_history", lambda db, user_id: history)
    embedded_queries = []
    monkeypatch.setattr(
        ai,
        "generate_embedding",
        lambda prompt: embedded_queries.append(prompt) or [0.0] * 1536,
    )
    monkeypatch.setattr(
        crud,
        "search_similar_chunks",
        lambda db, embedding, user_id, document_id=None: [SimpleNamespace(content="retrieved context")],
    )
    generated = {}

    def fake_generate_response(prompt, context=None, history=None):
        generated["history"] = history
        return "Hello from test"

    monkeypatch.setattr(ai, "generate_response", fake_generate_response)
    saved = {}

    def fake_create_chat_message(db, user_id, prompt, response):
        saved["user_id"] = user_id

    monkeypatch.setattr(crud, "create_chat_message", fake_create_chat_message)

    response = client.post("/chat", json={"prompt": "Hi"})

    assert response.status_code == 200
    assert response.json()["reply"] == "Hello from test"
    assert saved["user_id"] == 1
    assert embedded_queries == ["Summarize my experience.\nHi"]
    assert generated["history"] == history


def test_chat_searches_only_the_selected_document(monkeypatch, authenticated_user):
    monkeypatch.setattr(crud, "get_recent_chat_history", lambda db, user_id: [])
    monkeypatch.setattr(ai, "generate_embedding", lambda prompt: [0.0] * 768)
    searched = {}

    def fake_search(db, embedding, user_id, document_id=None):
        searched["document_id"] = document_id
        return [SimpleNamespace(content="CV context")]

    monkeypatch.setattr(crud, "search_similar_chunks", fake_search)
    monkeypatch.setattr(ai, "generate_response", lambda prompt, context=None, history=None: "Answer")
    monkeypatch.setattr(crud, "create_chat_message", lambda *args: None)

    response = client.post("/chat", json={"prompt": "Summarize this", "document_id": 7})

    assert response.status_code == 200
    assert searched["document_id"] == 7


def test_retrieval_query_keeps_previous_question_only_for_short_follow_ups():
    history = [
        {"role": "user", "content": "What was the first project?"},
        {"role": "assistant", "content": "A payment platform."},
    ]

    assert main._retrieval_query("Tell me more", history) == "What was the first project?\nTell me more"
    assert main._retrieval_query("Which databases did you use at the bank?", history) == (
        "Which databases did you use at the bank?"
    )
    assert main._retrieval_query("Tell me more", []) == "Tell me more"


def test_chat_explains_when_account_has_no_indexed_chunks(monkeypatch, authenticated_user):
    monkeypatch.setattr(crud, "get_recent_chat_history", lambda db, user_id: [])
    monkeypatch.setattr(ai, "generate_embedding", lambda prompt: [0.0] * 768)
    monkeypatch.setattr(crud, "search_similar_chunks", lambda db, embedding, user_id, document_id=None: [])
    monkeypatch.setattr(
        ai,
        "generate_response",
        lambda prompt, context=None, history=None: pytest.fail("LLM should not be called without retrieved context"),
    )
    monkeypatch.setattr(crud, "create_chat_message", lambda *args: None)

    response = client.post("/chat", json={"prompt": "Summarize my CV"})

    assert response.status_code == 200
    assert "Upload a document" in response.json()["reply"]


def test_chat_returns_gateway_timeout_when_ollama_times_out(monkeypatch, authenticated_user):
    monkeypatch.setattr(crud, "get_recent_chat_history", lambda db, user_id: [])
    monkeypatch.setattr(ai, "generate_embedding", lambda prompt: [0.0] * 768)
    monkeypatch.setattr(
        crud,
        "search_similar_chunks",
        lambda db, embedding, user_id, document_id=None: [SimpleNamespace(content="CV context")],
    )

    def raise_timeout(prompt, context=None, history=None):
        raise ai.OllamaTimeoutError("Ollama timed out")

    monkeypatch.setattr(ai, "generate_response", raise_timeout)
    response = client.post("/chat", json={"prompt": "Summarize my CV"})

    assert response.status_code == 504
    assert response.json()["detail"] == "Ollama timed out"


def test_upload_docx_extracts_text_from_paragraphs_and_tables():
    docx = Document()
    docx.add_paragraph("Python developer")
    table = docx.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "FastAPI"
    file_content = BytesIO()
    docx.save(file_content)

    text = extract_document_text("resume.docx", file_content.getvalue())

    assert "Python developer" in text
    assert "FastAPI" in text


def test_upload_pdf_uses_ocr_for_pages_without_selectable_text(monkeypatch):
    import pymupdf

    from PIL import Image

    pdf = pymupdf.open()
    page = pdf.new_page()
    image_bytes = BytesIO()
    Image.new("RGB", (32, 32), "white").save(image_bytes, format="PNG")
    page.insert_image(pymupdf.Rect(0, 0, 100, 100), stream=image_bytes.getvalue())
    file_content = pdf.tobytes()
    pdf.close()
    ocr_calls = []

    monkeypatch.setattr(
        "app.document_loader._ocr_pdf_page",
        lambda page: ocr_calls.append(page.number) or "Recognized text from scan",
    )

    text = extract_document_text("scanned.pdf", file_content)

    assert text == "Recognized text from scan"
    assert ocr_calls == [0]


def test_upload_pdf_does_not_require_ocr_for_selectable_text():
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Selectable PDF text")
    file_content = pdf.tobytes()
    pdf.close()

    text = extract_document_text("text.pdf", file_content)

    assert "Selectable PDF text" in text


def test_upload_pdf_ocr_includes_images_on_pages_with_selectable_text(monkeypatch):
    import pymupdf

    from PIL import Image

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_text((72, 72), "Selectable PDF text")
    image_bytes = BytesIO()
    Image.new("RGB", (32, 32), "white").save(image_bytes, format="PNG")
    page.insert_image(pymupdf.Rect(0, 0, 100, 100), stream=image_bytes.getvalue())
    file_content = pdf.tobytes()
    pdf.close()
    ocr_calls = []

    monkeypatch.setattr(
        "app.document_loader._ocr_pdf_image",
        lambda document, image_xref: ocr_calls.append(image_xref) or "Text in embedded image",
    )

    text = extract_document_text("mixed.pdf", file_content)

    assert "Selectable PDF text" in text
    assert "Text in embedded image" in text
    assert ocr_calls


def test_upload_pdf_ocrs_a_repeated_image_only_once(monkeypatch):
    import pymupdf

    from PIL import Image

    pdf = pymupdf.open()
    image_bytes = BytesIO()
    Image.new("RGB", (32, 32), "white").save(image_bytes, format="PNG")
    for _ in range(3):
        page = pdf.new_page()
        page.insert_text((72, 72), "Selectable PDF text")
        page.insert_image(pymupdf.Rect(0, 0, 100, 100), stream=image_bytes.getvalue())
    file_content = pdf.tobytes(garbage=4)
    pdf.close()
    ocr_calls = []

    monkeypatch.setattr(
        "app.document_loader._ocr_pdf_image",
        lambda document, image_xref: ocr_calls.append(image_xref) or "Logo text",
    )

    extract_document_text("logo.pdf", file_content)

    assert len(ocr_calls) == 1


def test_upload_document_file_indexes_extracted_text(monkeypatch, authenticated_user):
    docx = Document()
    docx.add_paragraph("CV text")
    file_content = BytesIO()
    docx.save(file_content)
    stored = {}

    def fake_store(db, title, content, source_url, user_id):
        stored.update(title=title, content=content, user_id=user_id)
        return {
            "id": 42,
            "title": title,
            "source_url": source_url,
            "chunk_count": 1,
        }

    monkeypatch.setattr(
        main,
        "_store_document",
        fake_store,
    )

    response = client.post(
        "/documents/upload",
        files={
            "file": (
                "resume.docx",
                file_content.getvalue(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )

    assert response.status_code == 201
    assert stored == {"title": "resume", "content": "CV text", "user_id": 1}
    assert response.json() == {
        "id": 42,
        "title": "resume",
        "source_url": None,
        "chunk_count": 1,
    }


def test_upload_document_file_rejects_unsupported_type(authenticated_user):
    response = client.post(
        "/documents/upload",
        files={"file": ("resume.txt", b"CV text", "text/plain")},
    )

    assert response.status_code == 422
    assert response.json()["detail"] == "Unsupported file type. Upload a PDF or DOCX file."


def test_upload_document_file_requires_authentication():
    docx = Document()
    docx.add_paragraph("CV text")
    file_content = BytesIO()
    docx.save(file_content)

    response = client.post(
        "/documents/upload",
        files={"file": ("resume.docx", file_content.getvalue(), "application/octet-stream")},
    )

    assert response.status_code == 401
