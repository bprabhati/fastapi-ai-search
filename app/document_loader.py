from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
import pymupdf
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from PIL import Image
import pytesseract


class OCRUnavailableError(RuntimeError):
    pass


def _ocr_pdf_page(page: pymupdf.Page) -> str:
    pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
    image = Image.open(BytesIO(pixmap.tobytes("png")))
    return _recognize_image_text(image)


def _ocr_pdf_image(document: pymupdf.Document, image_xref: int) -> str:
    pixmap = pymupdf.Pixmap(document, image_xref)
    if pixmap.n - pixmap.alpha > 3:
        pixmap = pymupdf.Pixmap(pymupdf.csRGB, pixmap)
    image = Image.open(BytesIO(pixmap.tobytes("png")))
    return _recognize_image_text(image)


def _recognize_image_text(image: Image.Image) -> str:
    try:
        return pytesseract.image_to_string(image, lang="eng")
    except pytesseract.TesseractNotFoundError as exc:
        raise OCRUnavailableError(
            "OCR is required for this PDF, but Tesseract is not installed. "
            "Use the Docker setup or install Tesseract OCR locally."
        ) from exc


def extract_document_text(filename: str, file_content: bytes) -> str:
    extension = PurePosixPath(filename).suffix.lower()

    if extension == ".pdf":
        ocr_document = None
        try:
            reader = PdfReader(BytesIO(file_content))
            ocr_document = pymupdf.open(stream=file_content, filetype="pdf")
            pages = []
            ocr_image_xrefs = set()  # a logo repeated on every page only needs OCR once
            for page_number, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                pymupdf_page = ocr_document[page_number]
                if not text.strip():
                    text = _ocr_pdf_page(pymupdf_page)
                else:
                    page_images = pymupdf_page.get_images(full=True)
                    if page_images:
                        new_xrefs = {image[0] for image in page_images if image[0]} - ocr_image_xrefs
                        ocr_image_xrefs |= new_xrefs
                        image_text = [_ocr_pdf_image(ocr_document, xref) for xref in sorted(new_xrefs)]
                        text = "\n".join([text, *image_text])
                pages.append(text)
            return "\n".join(pages)
        except (PdfReadError, OSError, ValueError) as exc:
            raise ValueError("The PDF could not be read. Check that it is a valid, unencrypted PDF.") from exc
        finally:
            if ocr_document is not None:
                ocr_document.close()

    if extension == ".docx":
        try:
            document = Document(BytesIO(file_content))
        except (PackageNotFoundError, BadZipFile, KeyError, ValueError) as exc:
            raise ValueError("The DOCX file could not be read. Check that it is a valid DOCX document.") from exc

        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_rows = [
            " ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        return "\n".join(text for text in paragraphs + table_rows if text.strip())

    raise ValueError("Unsupported file type. Upload a PDF or DOCX file.")
